"""
Academic Documentation Generator for Device Risk Analysis
Generates both DOCX and PDF formats for:
1. BRD_Device_Risk_Analysis (Business Requirements Document)
2. SRS_Device_Risk_Analysis (Software Requirements Specification)
3. SOW_Device_Risk_Analysis (Statement of Work)
4. WBS_Device_Risk_Analysis (Work Breakdown Structure)
5. Test_Cases_Device_Risk_Analysis (Comprehensive Test Suite)
6. Project_Report_Device_Risk_Analysis (Comprehensive Academic Project Report)
7. Presentation_Script_Device_Risk_Analysis (Viva Presentation Guide & Model Answers)

Author: Atharva Gahine (B.Tech CSE Semester V, ITM SKILLS UNIVERSITY)
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY, TA_RIGHT

os.makedirs("docs", exist_ok=True)

# -----------------------------------------------------------------------------
# DOCX STYLING UTILITIES
# -----------------------------------------------------------------------------
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_styled_docx():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
    return doc

def add_docx_cover(doc, doc_title, doc_subtitle):
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni = p_uni.add_run("ITM SKILLS UNIVERSITY\\nSchool of Future Tech\\nDepartment of Computer Science & Engineering\\n")
    r_uni.font.name = 'Helvetica'
    r_uni.font.size = Pt(13)
    r_uni.font.bold = True
    r_uni.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph().paragraph_format.space_after = Pt(24)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(doc_title)
    r_title.font.name = 'Helvetica'
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(37, 99, 235)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run(doc_subtitle)
    r_sub.font.name = 'Helvetica'
    r_sub.font.size = Pt(14)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph().paragraph_format.space_after = Pt(40)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run(
        "Candidate Name: Atharva Gahine\\n"
        "Degree: Bachelor of Technology (Computer Science & Engineering)\\n"
        "Semester: V | Academic Session: 2024–2028\\n"
        "Course: Machine Learning (Case Study No. 86)\\n"
        "Submission Date: October 2026"
    )
    r_meta.font.name = 'Helvetica'
    r_meta.font.size = Pt(11)
    r_meta.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_page_break()

def add_heading_1(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    r = h.add_run(text)
    r.font.name = 'Helvetica'
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)

def add_heading_2(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    r = h.add_run(text)
    r.font.name = 'Helvetica'
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(37, 99, 235)

def add_para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Helvetica'
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
    r = p.add_run(text)
    r.font.name = 'Helvetica'
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(51, 65, 85)

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Helvetica'
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
    r = p.add_run(text)
    r.font.name = 'Helvetica'
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(51, 65, 85)

def add_table_data(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    # Header row
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], '1E293B')
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = 'Helvetica'
                run.font.size = Pt(9.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)

    # Data rows
    for r_idx, row in enumerate(data):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_col = 'F8FAFC' if r_idx % 2 == 0 else 'FFFFFF'
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_col)
            for p in row_cells[c_idx].paragraphs:
                for run in p.runs:
                    run.font.name = 'Helvetica'
                    run.font.size = Pt(9.0)
                    run.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

# -----------------------------------------------------------------------------
# REPORTLAB PDF STYLING UTILITIES
# -----------------------------------------------------------------------------
def get_pdf_styles():
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#1E3A8A'),
        alignment=TA_CENTER,
        spaceAfter=10
    )
    subtitle_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#475569'),
        alignment=TA_CENTER,
        spaceAfter=25
    )
    meta_style = ParagraphStyle(
        'DocMeta',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#334155'),
        alignment=TA_CENTER,
        spaceAfter=30
    )
    h1_style = ParagraphStyle(
        'Heading1Custom',
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#0F172A'),
        spaceBefore=14,
        spaceAfter=6
    )
    h2_style = ParagraphStyle(
        'Heading2Custom',
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor('#2563EB'),
        spaceBefore=10,
        spaceAfter=4
    )
    body_style = ParagraphStyle(
        'BodyCustom',
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#334155'),
        alignment=TA_JUSTIFY,
        spaceAfter=6
    )
    bullet_style = ParagraphStyle(
        'BulletCustom',
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#334155'),
        leftIndent=15,
        spaceAfter=3
    )
    table_cell_style = ParagraphStyle(
        'TableCell',
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#1E293B')
    )
    table_hdr_style = ParagraphStyle(
        'TableHdr',
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.white,
        alignment=TA_CENTER
    )
    return {
        'title': title_style, 'subtitle': subtitle_style, 'meta': meta_style,
        'h1': h1_style, 'h2': h2_style, 'body': body_style, 'bullet': bullet_style,
        'cell': table_cell_style, 'hdr': table_hdr_style
    }

def build_pdf_doc(filename, elements):
    pdf = SimpleDocTemplate(
        filename,
        pagesize=letter,
        leftMargin=54, rightMargin=54, topMargin=54, bottomMargin=54
    )
    pdf.build(elements)

print("Documentation builder module loaded.")
