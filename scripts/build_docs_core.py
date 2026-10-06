"""
Comprehensive Academic Documentation Generator for Device Risk Analysis
Generates both .docx and .pdf files for:
1. BRD (Business Requirements Document)
2. SRS (Software Requirements Specification)
3. SOW (Statement of Work)
4. WBS (Work Breakdown Structure)
5. Test Cases Document
6. Project Report (Comprehensive Academic Project Report)
7. Presentation Script & Viva Preparation Guide

Author: Atharva Gahine (B.Tech CSE Semester V, ITM SKILLS UNIVERSITY)
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY, TA_RIGHT

os.makedirs("docs", exist_ok=True)

# -----------------------------------------------------------------------------
# DOCX HELPERS
# -----------------------------------------------------------------------------
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_styled_docx():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
    return doc

def add_docx_cover(doc, doc_title, doc_subtitle):
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni = p_uni.add_run("ITM SKILLS UNIVERSITY\nSchool of Future Tech\nDepartment of Computer Science & Engineering\n")
    r_uni.font.name = 'Helvetica'
    r_uni.font.size = Pt(13)
    r_uni.font.bold = True
    r_uni.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph().paragraph_format.space_after = Pt(24)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(doc_title)
    r_title.font.name = 'Helvetica'
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(37, 99, 235)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run(doc_subtitle)
    r_sub.font.name = 'Helvetica'
    r_sub.font.size = Pt(13)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph().paragraph_format.space_after = Pt(40)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run(
        "Candidate Name: Atharva Gahine\n"
        "Degree: Bachelor of Technology (Computer Science & Engineering)\n"
        "Semester: V | Academic Session: 2024–2028\n"
        "Course: Machine Learning (Case Study No. 86)\n"
        "Submission Date: October 2026"
    )
    r_meta.font.name = 'Helvetica'
    r_meta.font.size = Pt(11)
    r_meta.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_page_break()

def add_h1(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    r = h.add_run(text)
    r.font.name = 'Helvetica'
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)

def add_h2(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(11)
    h.paragraph_format.space_after = Pt(4)
    r = h.add_run(text)
    r.font.name = 'Helvetica'
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(37, 99, 235)

def add_p(doc, text, bold_pre=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    if bold_pre:
        r_pre = p.add_run(bold_pre)
        r_pre.font.name = 'Helvetica'
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
    r = p.add_run(text)
    r.font.name = 'Helvetica'
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(51, 65, 85)

def add_bullet(doc, text, bold_pre=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_pre:
        r_pre = p.add_run(bold_pre)
        r_pre.font.name = 'Helvetica'
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(30, 41, 59)
    r = p.add_run(text)
    r.font.name = 'Helvetica'
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(51, 65, 85)

def add_table(doc, headers, data):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], '1E293B')
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = 'Helvetica'
                run.font.size = Pt(9.0)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row in enumerate(data):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_col = 'F8FAFC' if r_idx % 2 == 0 else 'FFFFFF'
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_col)
            for p in row_cells[c_idx].paragraphs:
                for run in p.runs:
                    run.font.name = 'Helvetica'
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

# -----------------------------------------------------------------------------
# REPORTLAB HELPERS
# -----------------------------------------------------------------------------
styles = getSampleStyleSheet()
pdf_title = ParagraphStyle('DocTitle', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=18, leading=22, textColor=colors.HexColor('#1E3A8A'), alignment=TA_CENTER, spaceAfter=8)
pdf_sub = ParagraphStyle('DocSub', parent=styles['Normal'], fontName='Helvetica-Oblique', fontSize=11, leading=15, textColor=colors.HexColor('#475569'), alignment=TA_CENTER, spaceAfter=16)
pdf_meta = ParagraphStyle('DocMeta', parent=styles['Normal'], fontName='Helvetica', fontSize=9, leading=13, textColor=colors.HexColor('#334155'), alignment=TA_CENTER, spaceAfter=22)
pdf_h1 = ParagraphStyle('H1', fontName='Helvetica-Bold', fontSize=13, leading=17, textColor=colors.HexColor('#0F172A'), spaceBefore=12, spaceAfter=5)
pdf_h2 = ParagraphStyle('H2', fontName='Helvetica-Bold', fontSize=10.5, leading=14, textColor=colors.HexColor('#2563EB'), spaceBefore=8, spaceAfter=3)
pdf_p = ParagraphStyle('Body', fontName='Helvetica', fontSize=8.5, leading=12, textColor=colors.HexColor('#334155'), alignment=TA_JUSTIFY, spaceAfter=5)
pdf_bullet = ParagraphStyle('Bullet', fontName='Helvetica', fontSize=8.5, leading=12, textColor=colors.HexColor('#334155'), leftIndent=12, spaceAfter=2)
pdf_cell = ParagraphStyle('Cell', fontName='Helvetica', fontSize=7.5, leading=9.5, textColor=colors.HexColor('#1E293B'))
pdf_hdr = ParagraphStyle('Hdr', fontName='Helvetica-Bold', fontSize=8, leading=10, textColor=colors.white, alignment=TA_CENTER)

def make_pdf_table(headers, data, col_widths=None):
    table_data = [[Paragraph(h, pdf_hdr) for h in headers]]
    for row in data:
        table_data.append([Paragraph(str(val), pdf_cell) for val in row])
    t = Table(table_data, colWidths=col_widths)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#F8FAFC'), colors.white]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
    ]))
    return t

print("Helper definitions initialized.")
